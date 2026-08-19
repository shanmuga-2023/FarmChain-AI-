import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import matplotlib.pyplot as plt
import matplotlib.patches as patches

def create_architecture_diagram():
    fig, ax = plt.subplots(figsize=(7.6, 4.3), dpi=300)
    fig.patch.set_facecolor('#111318')
    ax.set_facecolor('#111318')

    # Color palette
    c_red = '#D91E27'
    c_blue = '#3B82F6'
    c_green = '#10B981'
    c_amber = '#F59E0B'
    c_card = '#1C1F2B'
    c_subcard = '#282D3F'
    c_white = '#FFFFFF'
    c_gray = '#94A3B8'
    c_light = '#E2E8F0'

    # Layer 1: Client & Presentation Layer
    r1 = patches.FancyBboxPatch((0.03, 0.73), 0.94, 0.23, boxstyle='round,pad=0.015,rounding_size=0.03',
                               facecolor=c_card, edgecolor=c_red, linewidth=1.8)
    ax.add_patch(r1)
    ax.text(0.5, 0.91, 'PRESENTATION & STAKEHOLDER LAYER (Vite + Vanilla JS SPA)', 
            ha='center', va='center', color=c_white, fontsize=8.5, fontweight='bold', family='sans-serif')

    boxes_l1 = [
        ('Farmer App\n(Voice AI: HI, TA, EN)', 0.15),
        ('Intermediary\n(Wholesale & Logistics)', 0.38),
        ('Retailer Portal\n(Shelf QR & Sourcing)', 0.62),
        ('Consumer & Admin\n(QR Trace & Security)', 0.85)
    ]
    for text, xpos in boxes_l1:
        sub = patches.FancyBboxPatch((xpos - 0.10, 0.75), 0.20, 0.12, boxstyle='round,pad=0.01,rounding_size=0.02',
                                     facecolor=c_subcard, edgecolor='#3E455E', linewidth=1)
        ax.add_patch(sub)
        ax.text(xpos, 0.81, text, ha='center', va='center', color=c_light, fontsize=7.0, fontweight='bold', family='sans-serif')

    # Connecting Arrow 1 -> 2
    ax.annotate('', xy=(0.5, 0.67), xytext=(0.5, 0.73),
                arrowprops=dict(arrowstyle='->,head_width=0.4,head_length=0.6', color=c_red, lw=2))

    # Layer 2: API & AI Engine
    r2_1 = patches.FancyBboxPatch((0.03, 0.38), 0.45, 0.27, boxstyle='round,pad=0.015,rounding_size=0.03',
                                  facecolor=c_card, edgecolor=c_blue, linewidth=1.5)
    ax.add_patch(r2_1)
    ax.text(0.255, 0.60, 'REAL-TIME BACKEND & ORACLE', ha='center', va='center', color='#60A5FA', fontsize=8, fontweight='bold', family='sans-serif')
    
    t_api = "• Node.js & Express.js REST APIs\n• Socket.io 2-Event Sync Service\n• eNAM & AgMarkNet Mandi Oracle\n• Local Resilient Offline Rate Cache"
    ax.text(0.06, 0.48, t_api, ha='left', va='center', color=c_light, fontsize=6.8, family='sans-serif', linespacing=1.4)

    r2_2 = patches.FancyBboxPatch((0.52, 0.38), 0.45, 0.27, boxstyle='round,pad=0.015,rounding_size=0.03',
                                  facecolor=c_card, edgecolor=c_green, linewidth=1.5)
    ax.add_patch(r2_2)
    ax.text(0.745, 0.60, 'AI & INTELLIGENCE ENGINE', ha='center', va='center', color='#34D399', fontsize=8, fontweight='bold', family='sans-serif')
    
    t_ai = "• APMC Fair Price Benchmark ML\n• Real-Time Anomaly & Fraud Detector\n• Predictive Demand Forecasting\n• Vernacular Speech NLP Parser"
    ax.text(0.55, 0.48, t_ai, ha='left', va='center', color=c_light, fontsize=6.8, family='sans-serif', linespacing=1.4)

    # Connecting Arrows 2 -> 3
    ax.annotate('', xy=(0.255, 0.31), xytext=(0.255, 0.38),
                arrowprops=dict(arrowstyle='->,head_width=0.4,head_length=0.6', color=c_blue, lw=1.8))
    ax.annotate('', xy=(0.745, 0.31), xytext=(0.745, 0.38),
                arrowprops=dict(arrowstyle='->,head_width=0.4,head_length=0.6', color=c_green, lw=1.8))

    # Layer 3: Blockchain & Smart Contracts
    r3 = patches.FancyBboxPatch((0.03, 0.04), 0.94, 0.25, boxstyle='round,pad=0.015,rounding_size=0.03',
                               facecolor=c_card, edgecolor=c_amber, linewidth=1.8)
    ax.add_patch(r3)
    ax.text(0.5, 0.24, 'TRUSTLESS BLOCKCHAIN & SOLIDITY SMART CONTRACT SUITE (EVM / Sepolia)', 
            ha='center', va='center', color='#FBBF24', fontsize=8.2, fontweight='bold', family='sans-serif')

    boxes_l3 = [
        ('MarketplaceEscrow\n(60/20/15/5% Split)', 0.15),
        ('ProductRegistry\n(Immutable Batch QR)', 0.38),
        ('SupplyChainTracker\n(Cold-chain Logs)', 0.62),
        ('QualityCertifier\n(AGMARK & Organic)', 0.85)
    ]
    for text, xpos in boxes_l3:
        sub = patches.FancyBboxPatch((xpos - 0.10, 0.07), 0.20, 0.13, boxstyle='round,pad=0.01,rounding_size=0.02',
                                     facecolor=c_subcard, edgecolor='#3E455E', linewidth=1)
        ax.add_patch(sub)
        ax.text(xpos, 0.135, text, ha='center', va='center', color=c_light, fontsize=6.8, fontweight='bold', family='sans-serif')

    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    plt.subplots_adjust(left=0, right=1, top=1, bottom=0)
    plt.savefig('architecture_diagram.png', dpi=300, facecolor='#111318', edgecolor='none')
    plt.close()
    print("Generated architecture_diagram.png")

def format_text_box(shape, header, body_lines, header_color=RGBColor(217, 30, 39), header_size=Pt(12), body_size=Pt(9.5)):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    
    if header:
        p_head = tf.paragraphs[0]
        p_head.text = header
        p_head.font.bold = True
        p_head.font.size = header_size
        p_head.font.name = "Calibri"
        p_head.font.color.rgb = header_color
        p_head.space_after = Pt(3)
        p_head.line_spacing = 1.05
    
    first_body = True if not header else False
    for line in body_lines:
        if first_body:
            p = tf.paragraphs[0]
            first_body = False
        else:
            p = tf.add_paragraph()
        p.text = line
        p.font.size = body_size
        p.font.name = "Calibri"
        p.font.color.rgb = RGBColor(240, 240, 240)
        p.space_after = Pt(2)
        p.line_spacing = 1.05

create_architecture_diagram()
print("Ready to assemble presentation")
